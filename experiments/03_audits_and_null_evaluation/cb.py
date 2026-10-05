import os as _config_os
import re as _config_re

def _required(name):
    value = _config_os.environ.get(name)
    if value is None or not value.strip():
        raise RuntimeError(f'Set {name} before running this script')
    return value

def _required_int(name):
    value = int(_required(name))
    if value <= 0:
        raise ValueError(f'{name} must be a positive integer')
    return value

def _configured(value):
    defaults = {'BCFTOOLS': 'bcftools', 'PLINK2': 'plink2', 'PYTHON': 'python3', 'TABIX': 'tabix', 'BGZIP': 'bgzip', 'SAMTOOLS': 'samtools', 'BEDTOOLS': 'bedtools', 'CROSSMAP': 'CrossMap', 'UDOCKER_BIN': 'udocker'}
    def resolve(match):
        name = match.group(1) or match.group(2)
        if name in _config_os.environ:
            return _required(name)
        return defaults.get(name) or _required(name)
    return _config_re.sub(r"\$\{([A-Z][A-Z0-9_]*)\}|\$([A-Z][A-Z0-9_]*)", resolve, value)

def _source_path(base, value):
    return value if _config_os.path.isabs(value) else _config_os.path.join(base, value)
import glob, sys
try:
    import openpyxl
except ImportError:
    print('NO_OPENPYXL')
    sys.exit(0)
D = _configured('${METADATA_DIR}')
f = [_required('COHORT_C_METADATA_FILE')][0]
print('file:', f.split('/')[-1])
wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
print('sheets:', wb.sheetnames)
hits = 0
for ws in wb.worksheets:
    for row in ws.iter_rows(values_only=True):
        cells = [str(c) if c is not None else '' for c in row]
        joined = ' | '.join(cells)
        if 'TCHL' in joined.upper():
            print(f'[{ws.title}] {joined[:400]}')
            hits += 1
            if hits > 25:
                break
    if hits > 25:
        break
print('TCHL rows found:', hits)
